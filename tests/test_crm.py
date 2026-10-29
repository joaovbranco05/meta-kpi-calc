from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from io import BytesIO
from threading import Barrier

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError

from meta_kpi_calc.api.app import create_app
from meta_kpi_calc.core.config import Settings
from meta_kpi_calc.crm.api import _require_local_peer
from meta_kpi_calc.crm.import_files import ImportFileError, parse_meta_export
from meta_kpi_calc.crm.ingestion import LeadSubmission, change_lead_status, ingest_lead
from meta_kpi_calc.crm.models import Lead, LeadActivity, LeadActivityActor, LeadActivityType, LeadStatus


def test_xml_spreadsheet_parser_maps_meta_fields_and_treats_naive_time_as_utc() -> None:
    source = b'''<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet><Table>
    <Row><Cell><Data>id</Data></Cell><Cell><Data>created_time</Data></Cell><Cell><Data>phone_number</Data></Cell><Cell><Data>lead_status</Data></Cell></Row>
    <Row><Cell><Data>1234567890123456</Data></Cell><Cell><Data ss:Type="DateTime">2026-10-27T10:00:00</Data></Cell><Cell><Data ss:Type="Number">5511999999999</Data></Cell><Cell><Data>New</Data></Cell></Row>
    </Table></Worksheet></Workbook>'''
    source = source.replace(b"ss:Type", b"Type")
    submissions, report = parse_meta_export(source, "sample.xls")
    assert report.total_rows == 1
    assert report.rejected == 0
    assert submissions[0].meta_lead_id == "1234567890123456"
    assert submissions[0].values["meta_created_time"] == datetime(2026, 10, 27, 10, tzinfo=UTC)
    assert submissions[0].values["phone_number"] == "5511999999999"
    assert submissions[0].values["source_status"] == "New"


def test_parser_rejects_dynamic_columns_and_xml_entities_without_echoing_values() -> None:
    with pytest.raises(ImportFileError) as error:
        parse_meta_export(b"id,extra_question\n123,private answer\n", "leads.csv")
    assert error.value.code == "unknown_columns"
    attack = b'<!DOCTYPE x [<!ENTITY secret "synthetic-private">]><Workbook/>'
    with pytest.raises(ImportFileError) as error:
        parse_meta_export(attack, "bad.xls")
    assert error.value.code == "xml_declarations_not_allowed"
    assert "synthetic-private" not in str(error.value)


def test_parser_rejects_excess_xml_column_index_before_padding() -> None:
    source = b'''<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet><Table>
    <Row><Cell><Data>id</Data></Cell></Row>
    <Row><Cell><Data>synthetic</Data></Cell><Cell ss:Index="1000000"><Data>hidden</Data></Cell></Row>
    </Table></Worksheet></Workbook>'''
    with pytest.raises(ImportFileError) as error:
        parse_meta_export(source, "malformed.xls")
    assert error.value.code == "too_many_columns"


def test_parser_rejects_sequential_xml_columns_before_row_accumulation() -> None:
    cells = b"".join(b"<Cell><Data>id</Data></Cell>" for _ in range(101))
    source = (
        b'<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet">'
        b"<Worksheet><Table><Row>" + cells + b"</Row></Table></Worksheet></Workbook>"
    )
    with pytest.raises(ImportFileError) as error:
        parse_meta_export(source, "wide.xml")
    assert error.value.code == "too_many_columns"


def test_parser_rejects_nonempty_values_beyond_declared_csv_headers() -> None:
    submissions, report = parse_meta_export(
        b"id,full_name\nsynthetic-id,Synthetic,unmapped-payload\n", "malformed.csv"
    )
    assert submissions == []
    assert report.reasons == {"invalid_row": 1}


def test_csv_parser_validates_rows_and_reports_only_safe_reason_codes() -> None:
    submissions, report = parse_meta_export(
        b"id,created_time,is_organic\n123,not-a-date,true\n456,,false\n", "x.csv"
    )
    assert len(submissions) == 1
    assert report.total_rows == 2
    assert report.rejected == 1
    assert report.reasons == {"invalid_created_time": 1}
    assert submissions[0].values["is_organic"] is False


def test_xlsx_parser_reads_values_and_rejects_formula_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "full_name"])
    sheet.append(["synthetic-xlsx-id", "Synthetic"])
    output = BytesIO()
    workbook.save(output)
    parsed, report = parse_meta_export(output.getvalue(), "sample.xlsx")
    assert report.total_rows == 1
    assert parsed[0].values["full_name"] == "Synthetic"

    sheet["B2"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ImportFileError) as error:
        parse_meta_export(output.getvalue(), "formula.xlsx")
    assert error.value.code == "formulas_not_allowed"


def test_parser_rejects_xlsx_with_more_than_column_limit() -> None:
    workbook = Workbook()
    sheet = workbook.active
    for column in range(1, 102):
        sheet.cell(row=1, column=column, value="id" if column == 1 else f"unused_{column}")
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ImportFileError) as error:
        parse_meta_export(output.getvalue(), "wide.xlsx")
    assert error.value.code == "too_many_columns"


def test_idempotent_ingestion_preserves_internal_status_and_records_source_corrections(migrated_database) -> None:
    _, database = migrated_database
    submission = LeadSubmission("external-1234567890123456", {"full_name": "Synthetic", "source_status": "New"})
    with database.session_factory.begin() as session:
        first = ingest_lead(session, submission)
        lead_id = first.lead.id
        assert first.created
        change_lead_status(session, first.lead, LeadStatus.CONTACTED)
    with database.session_factory.begin() as session:
        replay = ingest_lead(session, LeadSubmission(submission.meta_lead_id, {"full_name": "Corrected", "source_status": "Other"}))
        assert not replay.created
        assert replay.lead.commercial_status == LeadStatus.CONTACTED
        assert replay.lead.full_name == "Corrected"
        activities = session.scalars(select(LeadActivity).where(LeadActivity.lead_id == lead_id).order_by(LeadActivity.id)).all()
        assert [item.activity_type for item in activities] == [LeadActivityType.IMPORTED, LeadActivityType.STATUS_CHANGED, LeadActivityType.SOURCE_CORRECTED]
        assert activities[0].actor_type == LeadActivityActor.SYSTEM_IMPORT
        assert activities[1].previous_status == LeadStatus.NEW
        assert activities[1].new_status == LeadStatus.CONTACTED
        assert activities[2].changed_fields == "full_name,source_status"
        assert activities[2].note is None


def test_meta_campaign_is_optional_and_duplicate_meta_id_is_a_database_constraint(migrated_database) -> None:
    _, database = migrated_database
    with database.session_factory.begin() as session:
        result = ingest_lead(session, LeadSubmission("meta-without-campaign", {"meta_campaign_id": "not-synced"}))
        assert result.lead.campaign_id is None
    with database.session_factory.begin() as session:
        session.add(Lead(meta_lead_id="meta-without-campaign"))
        with pytest.raises(IntegrityError):
            session.flush()
        session.rollback()


def test_rollback_keeps_status_and_activity_atomic(migrated_database) -> None:
    _, database = migrated_database
    with database.session_factory.begin() as session:
        lead = ingest_lead(session, LeadSubmission("rollback-synthetic-id", {})).lead
        lead_id = lead.id
    session = database.session_factory()
    try:
        lead = session.get(Lead, lead_id)
        change_lead_status(session, lead, LeadStatus.LOST)
        session.rollback()
    finally:
        session.close()
    with database.session_factory() as session:
        lead = session.get(Lead, lead_id)
        assert lead.commercial_status == LeadStatus.NEW
        assert session.scalar(select(func.count(LeadActivity.id)).where(LeadActivity.lead_id == lead_id)) == 1


def test_competing_replays_create_one_lead_and_one_initial_activity(migrated_database) -> None:
    _, database = migrated_database
    barrier = Barrier(2)

    def worker() -> None:
        with database.session_factory() as session:
            barrier.wait()
            try:
                ingest_lead(session, LeadSubmission("concurrent-synthetic-id", {"email": "synthetic@example.test"}))
                session.commit()
            except Exception:
                session.rollback()
                raise

    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(lambda _: worker(), range(2)))
    with database.session_factory() as session:
        assert session.scalar(select(func.count(Lead.id)).where(Lead.meta_lead_id == "concurrent-synthetic-id")) == 1
        assert session.scalar(select(func.count(LeadActivity.id))) == 1


def test_crm_api_lists_changes_status_and_rejects_nonlocal_peer(migrated_database) -> None:
    settings, database = migrated_database
    app = create_app(settings)
    app.state.database = database
    with database.session_factory.begin() as session:
        lead = ingest_lead(session, LeadSubmission("api-synthetic-id", {"full_name": "Synthetic"})).lead
        lead_id = lead.id
    with TestClient(app) as client:
        response = client.get("/api/crm/leads", params={"commercial_status": "NEW"})
        assert response.status_code == 200
        assert response.json()["total"] == 1
        changed = client.patch(f"/api/crm/leads/{lead_id}/status", json={"commercial_status": "LOST"})
        assert changed.status_code == 200
        assert changed.json()["commercial_status"] == "LOST"
        created = client.post(
            f"/api/crm/leads/{lead_id}/activities",
            json={"activity_type": "NOTE", "note": "Synthetic note"},
        )
        assert created.status_code == 201
        assert created.json()["actor_type"] == "LOCAL_OPERATOR"
        assert client.get(f"/api/crm/leads/{lead_id}/activities").json()[-1]["note"] == "Synthetic note"
        bad = client.post(f"/api/crm/leads/{lead_id}/activities", json={"activity_type": "IMPORTED"})
        assert bad.status_code == 422

    request = type("Request", (), {"client": type("Client", (), {"host": "192.0.2.1"})()})()
    with pytest.raises(Exception) as error:
        _require_local_peer(request)
    assert getattr(error.value, "status_code", None) == 403
