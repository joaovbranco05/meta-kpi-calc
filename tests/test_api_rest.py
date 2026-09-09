from datetime import UTC, date, datetime
from decimal import Decimal
from importlib.util import module_from_spec, spec_from_file_location
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from pydantic import SecretStr, ValidationError
from sqlalchemy import func, select

from meta_kpi_calc.api.app import create_app
from meta_kpi_calc.api.schemas import EnrollmentWrite
from meta_kpi_calc.core.config import Settings
from meta_kpi_calc.db.models import (
    Campaign,
    CampaignBrand,
    CampaignInsight,
    EnrollmentRecord,
    SyncStatus,
)
from meta_kpi_calc.db.session import Database
from meta_kpi_calc.services.meta_client import MetaClient, MetaClientError
from meta_kpi_calc.services.sync_service import (
    SyncResult,
    SyncService,
    SyncServiceError,
)

DAY_1 = date(2026, 9, 1)
DAY_2 = date(2026, 9, 2)
TOKEN = "api-token-sentinel"


def real_settings(base: Settings, **overrides: object) -> Settings:
    values: dict[str, object] = {
        "demo_mode": False,
        "database_url": base.database_url,
        "demo_database_url": base.demo_database_url,
        "meta_access_token": SecretStr(TOKEN),
        "meta_ad_account_id": "act_123",
    }
    values.update(overrides)
    return Settings(_env_file=None, **values)


def add_campaign(
    database: Database,
    meta_id: str,
    *,
    brand: CampaignBrand = CampaignBrand.NAO_CLASSIFICADA,
    course: str | None = None,
    effective_status: str = "ACTIVE",
    start_time: datetime | None = None,
) -> int:
    with database.session_factory.begin() as session:
        campaign = Campaign(
            meta_campaign_id=meta_id,
            name=f"Campaign {meta_id}",
            status="ACTIVE",
            effective_status=effective_status,
            objective="OUTCOME_LEADS",
            brand=brand,
            course=course,
            start_time=start_time,
        )
        session.add(campaign)
        session.flush()
        return campaign.id


def add_insight(database: Database, campaign_id: int, insight_date: date) -> int:
    with database.session_factory.begin() as session:
        insight = CampaignInsight(
            campaign_id=campaign_id,
            date_start=insight_date,
            date_stop=insight_date,
            spend=Decimal("12.34"),
            reach=100,
            impressions=120,
            clicks=11,
            inline_link_clicks=8,
            leads=3,
            qualified_leads=None,
            frequency=Decimal("1.200000"),
            meta_ctr=Decimal("6.666667"),
            meta_cpc=Decimal("1.542500"),
            meta_cpm=Decimal("102.833333"),
            raw_actions=[{"private": "must-not-leak"}],
            raw_cost_per_action_type=[{"private": "must-not-leak"}],
            raw_response={"private": "must-not-leak"},
            synced_at=datetime(2026, 9, 2, 12, tzinfo=UTC),
        )
        session.add(insight)
        session.flush()
        return insight.id


def enrollment_payload(campaign_id: int, **overrides: object) -> dict[str, object]:
    values: dict[str, object] = {
        "campaign_id": campaign_id,
        "reference_date": DAY_1.isoformat(),
        "course": "Administration",
        "contracted_enrollments": 3,
        "paying_enrollments": 2,
        "cancellations": 1,
        "expected_revenue": "1000.10",
        "received_revenue": "700.05",
        "contribution_margin": None,
        "notes": None,
    }
    values.update(overrides)
    return values


def test_connection_demo_and_incomplete_real_never_touch_http(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, _ = migrated_database

    def forbidden(_request: httpx.Request) -> httpx.Response:
        raise AssertionError("network-like transport must not be called")

    external = httpx.Client(transport=httpx.MockTransport(forbidden))
    try:
        demo = Settings(
            _env_file=None,
            database_url=base.database_url,
            demo_database_url=base.demo_database_url,
            meta_access_token=SecretStr(TOKEN),
            meta_ad_account_id="123",
        )
        with TestClient(create_app(demo, http_client=external)) as client:
            response = client.get("/api/meta/connection")
        assert response.json() == {
            "mode": "demo",
            "configured": True,
            "connected": False,
            "account": None,
        }

        incomplete = real_settings(base, meta_access_token=None)
        with TestClient(create_app(incomplete, http_client=external)) as client:
            response = client.get("/api/meta/connection")
        assert response.json() == {
            "mode": "real",
            "configured": False,
            "connected": False,
            "account": None,
        }
        assert external.is_closed is False
    finally:
        external.close()


def test_connection_real_uses_mock_and_sanitizes_failure(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, _ = migrated_database
    requests: list[httpx.Request] = []

    def transport(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        return httpx.Response(
            200,
            json={
                "id": "123",
                "name": "Account",
                "currency": "BRL",
                "timezone_name": "America/Sao_Paulo",
                "account_status": 1,
            },
        )

    external = httpx.Client(transport=httpx.MockTransport(transport))
    try:
        app = create_app(real_settings(base), http_client=external)
        with TestClient(app) as client:
            response = client.get("/api/meta/connection")
            assert response.status_code == 200
            assert response.json()["account"] == {
                "account_id": "123",
                "account_name": "Account",
                "currency": "BRL",
                "timezone_name": "America/Sao_Paulo",
                "account_status": 1,
            }
            app.state.meta_client = Mock(spec=MetaClient)
            app.state.meta_client.get_account.side_effect = MetaClientError(
                f"private {TOKEN} https://external.invalid/body"
            )
            failed = client.get("/api/meta/connection")
        assert len(requests) == 1
        assert failed.status_code == 502
        assert failed.json() == {
            "detail": {
                "code": "meta_connection_failed",
                "message": "Meta connection check failed.",
            }
        }
        assert TOKEN not in failed.text
        assert "external.invalid" not in failed.text
        assert external.is_closed is False
    finally:
        external.close()


def test_sync_maps_result_and_service_errors(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, _ = migrated_database
    app = create_app(base)
    service = Mock(spec=SyncService)
    service.run_sync.return_value = SyncResult(
        sync_run_id=7,
        status=SyncStatus.SUCCESS,
        date_start=DAY_1,
        date_stop=DAY_2,
        campaign_count=2,
        record_count=4,
        is_partial=False,
    )
    payload = {
        "date_start": DAY_1.isoformat(),
        "date_stop": DAY_2.isoformat(),
        "active_only": False,
        "meta_campaign_id": "campaign-1",
    }
    with TestClient(app) as client:
        app.state.sync_service = service
        response = client.post("/api/meta/sync", json=payload)
        assert response.status_code == 200
        assert response.json() == {
            "sync_run_id": 7,
            "status": "SUCCESS",
            "date_start": "2026-09-01",
            "date_stop": "2026-09-02",
            "campaign_count": 2,
            "record_count": 4,
            "is_partial": False,
        }
        service.run_sync.assert_called_once_with(
            DAY_1,
            DAY_2,
            active_only=False,
            meta_campaign_id="campaign-1",
        )

        for code, expected_status, expected_code in (
            ("sync_in_progress", 409, "sync_in_progress"),
            ("demo_mode", 503, "meta_sync_unavailable"),
            ("meta_not_configured", 503, "meta_sync_unavailable"),
            ("invalid_date_range", 422, "invalid_date_range"),
            ("future_date", 422, "future_date"),
            ("invalid_campaign_filter", 422, "invalid_campaign_filter"),
            ("sync_failed", 502, "meta_sync_failed"),
        ):
            service.run_sync.side_effect = SyncServiceError(
                f"private {TOKEN}", code=code, sync_run_id=9
            )
            failed = client.post("/api/meta/sync", json=payload)
            assert failed.status_code == expected_status
            assert failed.json()["detail"]["code"] == expected_code
            assert TOKEN not in failed.text
            if expected_code == "meta_sync_failed":
                assert failed.json()["detail"]["sync_run_id"] == 9


def test_sync_unavailable_and_request_validation_are_stable(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, _ = migrated_database
    with TestClient(create_app(base)) as client:
        unavailable = client.post(
            "/api/meta/sync",
            json={"date_start": "2026-09-01", "date_stop": "2026-09-01"},
        )
        invalid = client.post(
            "/api/meta/sync",
            json={"date_start": TOKEN, "date_stop": "2026-09-01", "extra": TOKEN},
        )
    assert unavailable.status_code == 503
    assert unavailable.json()["detail"]["code"] == "meta_sync_unavailable"
    assert invalid.status_code == 422
    assert invalid.json() == {
        "detail": {
            "code": "validation_error",
            "message": "Request validation failed.",
        }
    }
    assert TOKEN not in invalid.text


@pytest.mark.parametrize("active_only", ["false", "0", "off", 0, 1])
def test_sync_active_only_requires_json_boolean(
    migrated_database: tuple[Settings, Database], active_only: object
) -> None:
    base, _ = migrated_database
    with TestClient(create_app(base)) as client:
        response = client.post(
            "/api/meta/sync",
            json={
                "date_start": "2026-09-01",
                "date_stop": "2026-09-01",
                "active_only": active_only,
            },
        )
    assert response.status_code == 422
    assert response.json()["detail"] == {
        "code": "validation_error",
        "message": "Request validation failed.",
    }


def test_unexpected_failure_and_database_failure_are_sanitized(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, _ = migrated_database
    app = create_app(base)
    service = Mock(spec=SyncService)
    service.run_sync.side_effect = RuntimeError(f"private {TOKEN}")
    with TestClient(app, raise_server_exceptions=False) as client:
        app.state.sync_service = service
        unexpected = client.post(
            "/api/meta/sync",
            json={"date_start": "2026-09-01", "date_stop": "2026-09-01"},
        )
    assert unexpected.status_code == 500
    assert unexpected.json()["detail"]["code"] == "internal_error"
    assert TOKEN not in unexpected.text

    class BrokenDatabase:
        @staticmethod
        def session_factory() -> None:
            raise OSError(f"private {TOKEN}")

    broken_app = create_app(base)
    broken_app.state.database = BrokenDatabase()
    with TestClient(broken_app, raise_server_exceptions=False) as client:
        unavailable = client.get("/api/campaigns")
    assert unavailable.status_code == 503
    assert unavailable.json()["detail"] == {
        "code": "database_unavailable",
        "message": "Database request failed.",
    }
    assert TOKEN not in unavailable.text


def test_campaign_list_detail_filters_pagination_and_classification(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, database = migrated_database
    first_id = add_campaign(
        database,
        "meta-1",
        brand=CampaignBrand.RCTEC,
        start_time=datetime(2026, 9, 3, 10, 0, tzinfo=UTC),
    )
    second_id = add_campaign(
        database,
        "meta-2",
        brand=CampaignBrand.FECAF,
        course="Law",
        effective_status="PAUSED",
        start_time=datetime(2026, 9, 20, 10, 0, tzinfo=UTC),
    )

    with TestClient(create_app(base)) as client:
        page = client.get("/api/campaigns?offset=1&limit=1")
        assert page.status_code == 200
        assert page.json()["total"] == 2
        assert [item["id"] for item in page.json()["items"]] == [second_id]

        filtered = client.get(
            "/api/campaigns?brand=FECAF&effective_status=PAUSED&meta_campaign_id=meta-2"
        )
        assert [item["meta_campaign_id"] for item in filtered.json()["items"]] == [
            "meta-2"
        ]
        date_filtered = client.get(
            "/api/campaigns?date_start=2026-09-01&date_stop=2026-09-10"
        )
        assert [item["id"] for item in date_filtered.json()["items"]] == [first_id]
        assert client.get(f"/api/campaigns/{first_id}").json()["meta_campaign_id"] == "meta-1"
        assert client.get("/api/campaigns/99999").status_code == 404

        patched = client.patch(
            f"/api/campaigns/{first_id}/classification",
            json={"brand": "CURSO_COM_BOLSA", "course": "  Nursing  "},
        )
        assert patched.status_code == 200
        assert patched.json()["brand"] == "CURSO_COM_BOLSA"
        assert patched.json()["course"] == "Nursing"

        cleared = client.patch(
            f"/api/campaigns/{first_id}/classification", json={"course": None}
        )
        assert cleared.json()["brand"] == "CURSO_COM_BOLSA"
        assert cleared.json()["course"] is None

        for body in ({}, {"brand": None}, {"course": "  "}, {"unknown": 1}):
            invalid = client.patch(
                f"/api/campaigns/{first_id}/classification", json=body
            )
            assert invalid.status_code == 422
            assert invalid.json()["detail"]["code"] == "validation_error"
        assert client.patch(
            "/api/campaigns/99999/classification", json={"course": None}
        ).status_code == 404
        assert client.get("/api/campaigns?limit=101").status_code == 422


def test_campaign_insights_filter_order_decimal_and_no_raw(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, database = migrated_database
    campaign_id = add_campaign(database, "meta-insight")
    second_id = add_insight(database, campaign_id, DAY_2)
    first_id = add_insight(database, campaign_id, DAY_1)

    with TestClient(create_app(base)) as client:
        response = client.get(
            f"/api/campaigns/{campaign_id}/insights?date_start=2026-09-01&date_stop=2026-09-02"
        )
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == 2
        assert [item["id"] for item in body["items"]] == [first_id, second_id]
        item = body["items"][0]
        assert item["campaign_id"] == campaign_id
        assert item["meta_campaign_id"] == "meta-insight"
        assert item["spend"] == "12.34"
        assert item["qualified_leads"] is None
        rendered = response.text
        assert "raw_response" not in rendered
        assert "raw_actions" not in rendered
        assert "must-not-leak" not in rendered

        assert client.get(
            f"/api/campaigns/{campaign_id}/insights?date_start=2026-09-01"
        ).status_code == 422
        assert client.get(
            f"/api/campaigns/{campaign_id}/insights?date_start=2026-09-02&date_stop=2026-09-01"
        ).status_code == 422
        assert client.get("/api/campaigns/99999/insights").status_code == 404


def test_enrollment_create_list_replace_conflicts_and_rollback(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, database = migrated_database
    first_campaign = add_campaign(database, "meta-enrollment", brand=CampaignBrand.RCTEC)
    second_campaign = add_campaign(database, "meta-other", brand=CampaignBrand.FECAF)
    first_payload = enrollment_payload(first_campaign)

    with TestClient(create_app(base)) as client:
        created = client.post("/api/enrollments", json=first_payload)
        assert created.status_code == 201
        record_id = created.json()["id"]
        assert created.json()["meta_campaign_id"] == "meta-enrollment"
        assert created.json()["expected_revenue"] == "1000.10"
        assert created.json()["contribution_margin"] is None

        duplicate = client.post("/api/enrollments", json=first_payload)
        assert duplicate.status_code == 409
        assert duplicate.json()["detail"]["code"] == "duplicate_enrollment"
        assert "PUT" in duplicate.json()["detail"]["message"]

        created_second = client.post(
            "/api/enrollments",
            json=enrollment_payload(
                second_campaign,
                course="Law",
                reference_date=DAY_2.isoformat(),
            ),
        )
        assert created_second.status_code == 201

        listed = client.get(
            f"/api/enrollments?campaign_id={first_campaign}&brand=RCTEC&course=Administration&date_start=2026-09-01&date_stop=2026-09-01"
        )
        assert listed.status_code == 200
        assert listed.json()["total"] == 1
        assert listed.json()["items"][0]["id"] == record_id
        paged = client.get("/api/enrollments?offset=1&limit=1")
        assert paged.json()["total"] == 2
        assert paged.json()["items"][0]["id"] == created_second.json()["id"]

        replaced = client.put(
            f"/api/enrollments/{record_id}",
            json=enrollment_payload(
                first_campaign,
                reference_date=DAY_2.isoformat(),
                course="  Medicine  ",
                contracted_enrollments=0,
                paying_enrollments=3,
                cancellations=4,
                contribution_margin="120.25",
                notes="updated",
            ),
        )
        assert replaced.status_code == 200
        assert replaced.json()["course"] == "Medicine"
        assert replaced.json()["paying_enrollments"] == 3
        assert replaced.json()["cancellations"] == 4
        assert replaced.json()["contribution_margin"] == "120.25"

        conflict = client.put(
            f"/api/enrollments/{record_id}",
            json=enrollment_payload(
                second_campaign,
                course="Law",
                reference_date=DAY_2.isoformat(),
            ),
        )
        assert conflict.status_code == 409
        missing_campaign = client.post(
            "/api/enrollments", json=enrollment_payload(99999)
        )
        assert missing_campaign.status_code == 404
        assert client.put(
            "/api/enrollments/99999", json=enrollment_payload(first_campaign)
        ).status_code == 404
        assert client.get("/api/enrollments?date_start=2026-09-01").status_code == 422

    with database.session_factory() as session:
        enrollment = session.get(EnrollmentRecord, record_id)
        assert enrollment is not None
        assert enrollment.campaign_id == first_campaign
        assert enrollment.course == "Medicine"
        assert session.scalar(select(func.count(EnrollmentRecord.id))) == 2


def test_enrollment_body_validation_is_fixed_and_does_not_echo_input(
    migrated_database: tuple[Settings, Database],
) -> None:
    base, database = migrated_database
    campaign_id = add_campaign(database, "meta-validation")
    payload = enrollment_payload(campaign_id, course=TOKEN, expected_revenue="-1")
    with TestClient(create_app(base)) as client:
        response = client.post("/api/enrollments", json=payload)
    assert response.status_code == 422
    assert response.json()["detail"] == {
        "code": "validation_error",
        "message": "Request validation failed.",
    }
    assert TOKEN not in response.text


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("campaign_id", True),
        ("campaign_id", 1.0),
        ("campaign_id", "1"),
        ("contracted_enrollments", False),
        ("contracted_enrollments", 1.0),
        ("paying_enrollments", "2"),
        ("cancellations", False),
    ],
)
def test_enrollment_body_integers_are_strict(
    migrated_database: tuple[Settings, Database], field: str, value: object
) -> None:
    base, database = migrated_database
    campaign_id = add_campaign(database, f"strict-{field}-{type(value).__name__}")
    payload = enrollment_payload(campaign_id)
    payload[field] = value

    with TestClient(create_app(base)) as client:
        response = client.post("/api/enrollments", json=payload)

    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "validation_error"


def test_enrollment_money_becomes_finite_nonnegative_decimal() -> None:
    valid = enrollment_payload(1, expected_revenue=12.5, received_revenue="7.25")
    parsed = EnrollmentWrite.model_validate(valid)
    assert parsed.expected_revenue == Decimal("12.5")
    assert parsed.received_revenue == Decimal("7.25")
    assert isinstance(parsed.expected_revenue, Decimal)
    assert isinstance(parsed.received_revenue, Decimal)

    for field, value in (
        ("expected_revenue", True),
        ("received_revenue", "NaN"),
        ("expected_revenue", "Infinity"),
        ("contribution_margin", "-0.01"),
    ):
        invalid = enrollment_payload(1)
        invalid[field] = value
        with pytest.raises(ValidationError):
            EnrollmentWrite.model_validate(invalid)


def test_lifespan_owns_only_created_http_client_and_controls_scheduler(
    migrated_database: tuple[Settings, Database], monkeypatch: pytest.MonkeyPatch
) -> None:
    base, _ = migrated_database
    owned_client = Mock(spec=httpx.Client)
    scheduler = Mock()
    monkeypatch.setattr("meta_kpi_calc.api.app.httpx.Client", Mock(return_value=owned_client))
    create_scheduler_mock = Mock(return_value=scheduler)
    monkeypatch.setattr("meta_kpi_calc.api.app.create_scheduler", create_scheduler_mock)
    config = real_settings(base, scheduler_enabled=True)

    application = create_app(config)
    assert owned_client.close.call_count == 0
    with TestClient(application):
        scheduler.start.assert_called_once_with()
        scheduler.shutdown.assert_not_called()
    scheduler.shutdown.assert_called_once_with(wait=False)
    owned_client.close.assert_called_once_with()
    assert create_scheduler_mock.call_count == 1


def test_owned_client_closes_even_if_scheduler_shutdown_fails(
    migrated_database: tuple[Settings, Database], monkeypatch: pytest.MonkeyPatch
) -> None:
    base, _ = migrated_database
    owned_client = Mock(spec=httpx.Client)
    scheduler = Mock()
    scheduler.shutdown.side_effect = RuntimeError("shutdown failed")
    monkeypatch.setattr("meta_kpi_calc.api.app.httpx.Client", Mock(return_value=owned_client))
    monkeypatch.setattr(
        "meta_kpi_calc.api.app.create_scheduler", Mock(return_value=scheduler)
    )

    with pytest.raises(RuntimeError, match="shutdown failed"):
        with TestClient(create_app(real_settings(base, scheduler_enabled=True))):
            pass

    scheduler.start.assert_called_once_with()
    scheduler.shutdown.assert_called_once_with(wait=False)
    owned_client.close.assert_called_once_with()


def test_default_app_import_and_lifespan_do_not_create_external_client(
    migrated_database: tuple[Settings, Database], monkeypatch: pytest.MonkeyPatch
) -> None:
    base, _ = migrated_database
    http_constructor = Mock(side_effect=AssertionError("must not construct HTTP"))
    monkeypatch.setattr("meta_kpi_calc.api.app.httpx.Client", http_constructor)
    with TestClient(create_app(base)) as client:
        assert client.get("/health").status_code == 200
        assert client.get("/api/dashboard/summary").status_code == 422
        assert client.get("/api/export").status_code == 422
    http_constructor.assert_not_called()


def test_dashboard_and_exports_share_the_same_campaign_scope(
    migrated_database: tuple[Settings, Database],
) -> None:
    settings, database = migrated_database
    selected = add_campaign(
        database,
        "selected",
        brand=CampaignBrand.RCTEC,
        course="Administration",
        effective_status="ACTIVE",
    )
    other = add_campaign(
        database,
        "other",
        brand=CampaignBrand.RCTEC,
        course="Other",
        effective_status="ACTIVE",
    )
    add_insight(database, selected, DAY_1)
    add_insight(database, other, DAY_1)
    with database.session_factory.begin() as session:
        session.add(
            EnrollmentRecord(
                campaign_id=selected,
                reference_date=DAY_1,
                course="Administration",
                contracted_enrollments=1,
                paying_enrollments=1,
                cancellations=0,
                expected_revenue=Decimal("10.00"),
                received_revenue=Decimal("10.00"),
                contribution_margin=None,
                notes="=not-a-formula",
            )
        )

    params = {
        "date_start": DAY_1.isoformat(),
        "date_stop": DAY_1.isoformat(),
        "brand": "RCTEC",
        "course": "Administration",
        "effective_status": "ACTIVE",
    }
    with TestClient(create_app(settings)) as client:
        campaigns = client.get(
            "/api/campaigns",
            params={
                "brand": "RCTEC",
                "course": "Administration",
                "effective_status": "ACTIVE",
            },
        )
        assert campaigns.status_code == 200
        assert [item["id"] for item in campaigns.json()["items"]] == [selected]
        summary = client.get("/api/dashboard/summary", params=params)
        assert summary.status_code == 200
        body = summary.json()
        assert body["filters"]["course"] == "Administration"
        assert body["totals"]["spend"] == "12.34"
        assert body["totals"]["contracted_enrollments"] == 1

        csv_response = client.get(
            "/api/export", params={**params, "format": "enrollments", "dataset": "enrollments"}
        )
        assert csv_response.status_code == 422
        csv_response = client.get(
            "/api/export", params={**params, "format": "csv", "dataset": "enrollments"}
        )
        assert csv_response.status_code == 200
        assert csv_response.content.startswith(b"\xef\xbb\xbf")
        assert "raw_response" not in csv_response.text
        assert "'=not-a-formula" in csv_response.text
        assert csv_response.headers["content-disposition"] == (
            'attachment; filename="meta-kpi-enrollments.csv"'
        )

        workbook_response = client.get(
            "/api/export", params={**params, "format": "xlsx", "dataset": "performance"}
        )
    assert workbook_response.status_code == 200
    workbook = load_workbook(BytesIO(workbook_response.content), data_only=False)
    assert workbook.sheetnames == ["Resumo", "Diário", "Matrículas", "Campanhas"]
    enrollment_sheet = workbook["Matrículas"]
    headers = [cell.value for cell in enrollment_sheet[1]]
    notes_column = headers.index("notes") + 1
    assert enrollment_sheet.cell(row=2, column=notes_column).value == "'=not-a-formula"


def test_dashboard_and_export_reject_invalid_or_blank_filter_values(
    migrated_database: tuple[Settings, Database],
) -> None:
    settings, _ = migrated_database
    with TestClient(create_app(settings)) as client:
        for path, params in (
            ("/api/dashboard/summary", {"date_start": DAY_2, "date_stop": DAY_1}),
            ("/api/dashboard/summary", {"date_start": DAY_1, "date_stop": DAY_1, "course": "  "}),
            ("/api/export", {"date_start": DAY_1, "date_stop": DAY_1, "format": "csv", "dataset": "performance", "effective_status": "  "}),
        ):
            response = client.get(path, params=params)
            assert response.status_code == 422
            assert response.json()["detail"]["code"] == "validation_error"

        empty = client.get(
            "/api/dashboard/summary",
            params={"date_start": DAY_1, "date_stop": DAY_1, "course": "absent"},
        )
    assert empty.status_code == 200
    assert empty.json()["totals"]["spend"] == "0.00"
    assert empty.json()["warnings"] == []


def test_frontend_propagates_the_active_scope_to_each_supported_listing() -> None:
    path = Path(__file__).resolve().parents[1] / "frontend" / "app.py"
    spec = spec_from_file_location("frontend_app", path)
    assert spec is not None and spec.loader is not None
    frontend_app = module_from_spec(spec)
    spec.loader.exec_module(frontend_app)
    assert frontend_app.EFFECTIVE_STATUS_OPTIONS == (
        "",
        "ACTIVE",
        "PAUSED",
        "ARCHIVED",
        "DELETED",
        "DISABLED",
    )
    assert "FALSE" not in frontend_app.EFFECTIVE_STATUS_OPTIONS

    class FakeStreamlit:
        sidebar: "FakeStreamlit"

        def __init__(self, effective_status: str) -> None:
            self.sidebar = self
            self.effective_status = effective_status

        def __enter__(self) -> "FakeStreamlit":
            return self

        def __exit__(self, *_args: object) -> None:
            return None

        def header(self, _label: str) -> None:
            return None

        def date_input(self, label: str, _default: date | None = None) -> date:
            return DAY_1 if label == "Início" else DAY_2

        def selectbox(
            self, label: str, options: tuple[str, ...] | list[str], **kwargs: object
        ) -> str:
            if label == "Status efetivo":
                assert tuple(options) == frontend_app.EFFECTIVE_STATUS_OPTIONS
                format_func = kwargs["format_func"]
                assert callable(format_func)
                assert format_func("") == "Todos"
                assert format_func("ACTIVE") == "ACTIVE"
                assert format_func("DISABLED") == "DISABLED"
                return self.effective_status
            return ""

        def text_input(self, _label: str) -> str:
            return ""

    blank_filters = frontend_app._filters(FakeStreamlit(""))
    assert "effective_status" not in blank_filters
    active_filters = frontend_app._filters(FakeStreamlit("ACTIVE"))
    assert active_filters["effective_status"] == "ACTIVE"

    filters = {
        "date_start": DAY_1.isoformat(),
        "date_stop": DAY_2.isoformat(),
        "brand": "RCTEC",
        "campaign_id": "7",
        "course": "Administration",
        "effective_status": "ACTIVE",
    }
    assert frontend_app._campaign_params(filters) == {
        "limit": 100,
        "date_start": DAY_1.isoformat(),
        "date_stop": DAY_2.isoformat(),
        "brand": "RCTEC",
        "campaign_id": "7",
        "course": "Administration",
        "effective_status": "ACTIVE",
    }
    assert frontend_app._enrollment_params(filters) == {
        "limit": 100,
        "date_start": DAY_1.isoformat(),
        "date_stop": DAY_2.isoformat(),
        "brand": "RCTEC",
        "campaign_id": "7",
    }
