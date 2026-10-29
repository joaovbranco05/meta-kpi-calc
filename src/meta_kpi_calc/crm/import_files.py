"""Bounded parsers for Meta lead exports; parser errors never include source values."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, BinaryIO
from xml.etree import ElementTree

from openpyxl import load_workbook

from meta_kpi_calc.crm.ingestion import LeadSubmission

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 100
MAX_CELL_CHARS = 10_000
MAX_XLSX_ENTRIES = 200
MAX_XLSX_EXPANDED_BYTES = 100 * 1024 * 1024

EXPECTED_COLUMNS = {
    "id", "created_time", "ad_id", "ad_name", "adset_id", "adset_name",
    "campaign_id", "campaign_name", "form_id", "form_name", "is_organic",
    "platform", "full_name", "phone_number", "email", "city", "state", "lead_status",
}
REQUIRED_COLUMNS = {"id"}
COLUMN_MAP = {
    "id": "meta_lead_id", "created_time": "meta_created_time", "ad_id": "meta_ad_id",
    "ad_name": "meta_ad_name", "adset_id": "meta_adset_id", "adset_name": "meta_adset_name",
    "campaign_id": "meta_campaign_id", "campaign_name": "meta_campaign_name",
    "form_id": "meta_form_id", "form_name": "meta_form_name", "is_organic": "is_organic",
    "platform": "platform", "full_name": "full_name", "phone_number": "phone_number",
    "email": "email", "city": "city", "state": "state", "lead_status": "source_status",
}
FIELD_LIMITS = {
    "full_name": 255, "phone_number": 64, "email": 320, "city": 255,
    "state": 128, "meta_campaign_id": 64, "meta_campaign_name": 255,
    "meta_adset_id": 64, "meta_adset_name": 255, "meta_ad_id": 64,
    "meta_ad_name": 255, "meta_form_id": 64, "meta_form_name": 255,
    "source_status": 64, "platform": 64,
}


class ImportFileError(ValueError):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


@dataclass
class ImportReport:
    total_rows: int = 0
    inserted: int = 0
    existing: int = 0
    rejected: int = 0
    reasons: dict[str, int] = field(default_factory=dict)

    def reject(self, reason: str) -> None:
        self.rejected += 1
        self.reasons[reason] = self.reasons.get(reason, 0) + 1

    def as_dict(self) -> dict[str, Any]:
        return {
            "total_rows": self.total_rows,
            "inserted": self.inserted,
            "existing": self.existing,
            "rejected": self.rejected,
            "reasons": dict(sorted(self.reasons.items())),
        }


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _validate_headers(headers: list[str]) -> None:
    normalized = [header.strip() for header in headers]
    if not normalized or any(not item for item in normalized) or len(set(normalized)) != len(normalized):
        raise ImportFileError("invalid_headers")
    if len(normalized) > MAX_COLUMNS:
        raise ImportFileError("too_many_columns")
    if not REQUIRED_COLUMNS.issubset(normalized):
        raise ImportFileError("missing_required_columns")
    if set(normalized) - EXPECTED_COLUMNS:
        raise ImportFileError("unknown_columns")


def _csv_rows(data: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        text = data.decode("utf-8-sig")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample else csv.excel
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(io.StringIO(text), dialect)
        headers = next(rows)
        result = []
        for row in rows:
            result.append(row)
            if len(result) > MAX_ROWS:
                raise ImportFileError("too_many_rows")
        return headers, result
    except ImportFileError:
        raise
    except (UnicodeError, csv.Error, StopIteration) as exc:
        raise ImportFileError("invalid_csv") from exc


def _xlsx_rows(data: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ENTRIES or sum(item.file_size for item in infos) > MAX_XLSX_EXPANDED_BYTES:
                raise ImportFileError("xlsx_limits_exceeded")
            if any(item.file_size > MAX_XLSX_EXPANDED_BYTES for item in infos):
                raise ImportFileError("xlsx_limits_exceeded")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            sheet = workbook.active
            if sheet.max_column > MAX_COLUMNS:
                raise ImportFileError("too_many_columns")
            if sheet.max_row > MAX_ROWS + 1:
                raise ImportFileError("too_many_rows")
            iterator = sheet.iter_rows(values_only=False)
            header_cells = next(iterator)
            headers: list[str] = []
            for cell in header_cells:
                if len(headers) >= MAX_COLUMNS:
                    raise ImportFileError("too_many_columns")
                headers.append(str(cell.value or "").strip())
            values: list[list[Any]] = []
            for row in iterator:
                if any(cell.data_type == "f" for cell in row):
                    raise ImportFileError("formulas_not_allowed")
                row_values: list[Any] = []
                for cell in row:
                    if len(row_values) >= MAX_COLUMNS:
                        raise ImportFileError("too_many_columns")
                    row_values.append(cell.value)
                values.append(row_values)
                if len(values) > MAX_ROWS:
                    raise ImportFileError("too_many_rows")
            return headers, values
        finally:
            workbook.close()
    except ImportFileError:
        raise
    except (zipfile.BadZipFile, OSError, ValueError, StopIteration, KeyError) as exc:
        raise ImportFileError("invalid_xlsx") from exc


def _xml_rows(data: bytes) -> tuple[list[str], list[list[Any]]]:
    if re.search(br"<!\s*(DOCTYPE|ENTITY)\b", data, re.IGNORECASE):
        raise ImportFileError("xml_declarations_not_allowed")
    try:
        root = ElementTree.fromstring(data)
        worksheet = next((node for node in root.iter() if _local_name(node.tag) == "Worksheet"), None)
        table = next((node for node in worksheet.iter() if _local_name(node.tag) == "Table"), None) if worksheet is not None else None
        if table is None:
            raise ImportFileError("invalid_spreadsheet_xml")
        rows: list[list[Any]] = []
        for row_node in (node for node in table if _local_name(node.tag) == "Row"):
            if len(rows) >= MAX_ROWS + 1:
                raise ImportFileError("too_many_rows")
            row: list[Any] = []
            next_index = 1
            for cell in (node for node in row_node if _local_name(node.tag) == "Cell"):
                index = next((value for key, value in cell.attrib.items() if _local_name(key) == "Index"), None)
                if index is not None:
                    next_index = int(index)
                    if next_index < 1 or next_index > MAX_COLUMNS:
                        raise ImportFileError("too_many_columns")
                if next_index > MAX_COLUMNS:
                    raise ImportFileError("too_many_columns")
                while len(row) < next_index - 1:
                    if len(row) >= MAX_COLUMNS:
                        raise ImportFileError("too_many_columns")
                    row.append(None)
                data_node = next((node for node in cell if _local_name(node.tag) == "Data"), None)
                value: Any = data_node.text if data_node is not None else None
                dtype = next((value for key, value in (data_node.attrib.items() if data_node is not None else []) if _local_name(key) == "Type"), "String")
                if dtype == "DateTime" and value:
                    value = datetime.fromisoformat(value.replace("Z", "+00:00"))
                elif dtype == "Number" and value:
                    value = float(value) if "." in value else int(value)
                elif dtype == "Boolean" and value is not None:
                    value = value == "1"
                if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
                    raise ImportFileError("cell_too_long")
                if len(row) >= MAX_COLUMNS:
                    raise ImportFileError("too_many_columns")
                row.append(value)
                next_index = len(row) + 1
            rows.append(row)
        if not rows:
            raise ImportFileError("empty_file")
        return [str(value or "") for value in rows[0]], rows[1:]
    except ImportFileError:
        raise
    except (ElementTree.ParseError, ValueError, OverflowError) as exc:
        raise ImportFileError("invalid_spreadsheet_xml") from exc


def _scalar_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if abs(value) > 9_007_199_254_740_991:
            raise ImportFileError("invalid_scalar")
        if not value.is_integer():
            raise ImportFileError("invalid_scalar")
        return str(int(value))
    text = str(value).strip()
    return text or None


def _normalize_row(headers: list[str], row: list[Any]) -> LeadSubmission:
    if len(row) > len(headers) and any(value not in (None, "") for value in row[len(headers):]):
        raise ImportFileError("invalid_row")
    if len(row) > MAX_COLUMNS or any(isinstance(item, str) and len(item) > MAX_CELL_CHARS for item in row):
        raise ImportFileError("invalid_row")
    source = dict(zip(headers, row))
    if not isinstance(source.get("id"), str):
        raise ImportFileError("invalid_lead_id")
    lead_id = _scalar_text(source.get("id"))
    if not lead_id or len(lead_id) > 64:
        raise ImportFileError("invalid_lead_id")
    values: dict[str, Any] = {}
    for column, field_name in COLUMN_MAP.items():
        raw = source.get(column)
        if column == "created_time":
            if raw is None or raw == "":
                continue
            try:
                parsed = raw if isinstance(raw, datetime) else datetime.fromisoformat(str(raw).replace("Z", "+00:00"))
                values[field_name] = parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)
            except ValueError as exc:
                raise ImportFileError("invalid_created_time") from exc
        elif column == "is_organic":
            text = _scalar_text(raw)
            if text is None:
                continue
            lowered = text.lower()
            if lowered in {"true", "1"}:
                values[field_name] = True
            elif lowered in {"false", "0"}:
                values[field_name] = False
            else:
                raise ImportFileError("invalid_boolean")
        else:
            text = _scalar_text(raw)
            if text is not None:
                if column == "phone_number" and re.fullmatch(r"\d+\.0", text):
                    text = text[:-2]
                if len(text) > MAX_CELL_CHARS:
                    raise ImportFileError("cell_too_long")
                if len(text) > FIELD_LIMITS.get(field_name, MAX_CELL_CHARS):
                    raise ImportFileError("field_too_long")
                values[field_name] = text
    values["source"] = "META_LEAD_ADS"
    return LeadSubmission(lead_id, values)


def parse_meta_export(file: BinaryIO | bytes, filename: str = "") -> tuple[list[LeadSubmission], ImportReport]:
    data = file if isinstance(file, bytes) else file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError("file_too_large")
    if not data:
        raise ImportFileError("empty_file")
    stripped = data.lstrip()
    if re.search(br"<!\s*(DOCTYPE|ENTITY)\b", data[:65536], re.IGNORECASE):
        raise ImportFileError("xml_declarations_not_allowed")
    if stripped.startswith(b"<?xml") or b"urn:schemas-microsoft-com:office:spreadsheet" in data[:8192]:
        headers, rows = _xml_rows(data)
    elif data.startswith(b"PK\x03\x04"):
        headers, rows = _xlsx_rows(data)
    else:
        headers, rows = _csv_rows(data)
    headers = [str(value).strip() for value in headers]
    _validate_headers(headers)
    submissions: list[LeadSubmission] = []
    report = ImportReport()
    for row in rows:
        if not row or all(value is None or value == "" for value in row):
            continue
        report.total_rows += 1
        try:
            submissions.append(_normalize_row(headers, row))
        except ImportFileError as exc:
            report.reject(exc.code)
    return submissions, report
